"""Rendered regression for number cleanup, early breaks and unlabelled tails."""
import sys
import tempfile
from copy import deepcopy
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from guides.coursework_kfu_2025.table_continuation import (
    remove_single_page_column_numbers_inplace, refill_early_continuations_inplace,
    repair_remaining_table_spills_inplace, _merge_adjacent_numeric_tail,
    _split_cross_page_table_with_marker, _table_payload_rows,
    _cross_page_without_marker_probe,
)


def fixture(count=4, height=35, appendix=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    doc.styles['Normal'].font.size = Pt(12)
    if appendix:
        doc.add_paragraph('ПРИЛОЖЕНИЕ А')
    doc.add_paragraph('Таблица 1.1.1').paragraph_format.keep_with_next = True
    t = doc.add_table(rows=count + 2, cols=2)
    t.style = 'Table Grid'
    for i, row in enumerate(t.rows):
        values = ['Показатель', 'Описание'] if i == 0 else (
            ['1', '2'] if i == 1 else [f'Наблюдение номер {i}', f'Исходные данные этапа {i}'])
        for cell, text in zip(row.cells, values):
            cell.text = text
            cell.paragraphs[0].paragraph_format.keep_with_next = i < 2 or i == count + 1
        row.height = Pt(20 if i < 2 else height)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    doc.add_paragraph('Источник: исходные данные').paragraph_format.keep_with_next = True
    doc.add_paragraph('Примечание: проверенное пояснение')
    return doc


def check_clean(path, payload):
    assert _table_payload_rows(Document(path)) == payload
    fail, _, lines = _cross_page_without_marker_probe(path, None)
    assert not fail, fail
    return lines


def main():
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / 'trial.docx'
        for appendix in (False, True):
            d = fixture(1, appendix=appendix)
            payload = _table_payload_rows(d)
            d.save(p)
            assert remove_single_page_column_numbers_inplace(p) == 1
            assert len(Document(p).tables[0].rows) == 2
            check_clean(p, payload)
            before = p.read_bytes()
            assert remove_single_page_column_numbers_inplace(p) == 0 and p.read_bytes() == before
        print('PASS ordinary and captioned appendix: remove single-page numbers, preserve data, retry')
        d = fixture(15, 60)
        payload = _table_payload_rows(d)
        d.save(p)
        before = p.read_bytes()
        assert remove_single_page_column_numbers_inplace(p) == 0 and p.read_bytes() == before
        assert repair_remaining_table_spills_inplace(p) > 0
        check_clean(p, payload)
        before = p.read_bytes()
        assert remove_single_page_column_numbers_inplace(p) == 0 and p.read_bytes() == before
        print('PASS actual spills and labelled continuation numbers retained')
        d = fixture(10, 65)
        payload = _table_payload_rows(d)
        assert _split_cross_page_table_with_marker(d, 0, 4, '1.1.1', numeric_row_idx=1)
        d.save(p)
        assert refill_early_continuations_inplace(p) == 1
        lines = check_clean(p, payload)
        last_page = next(l.page_num for l in lines if 'Наблюдение номер 11' in l.text)
        assert all(any(l.page_num == last_page and text in l.text for l in lines)
                   for text in ['Источник:', 'Примечание:'])
        before = p.read_bytes()
        assert refill_early_continuations_inplace(p) == 0 and p.read_bytes() == before
        print('PASS early split refilled, final row and notes together, retry')
        d = fixture(10, 65)
        _split_cross_page_table_with_marker(d, 0, 4, '1.1.1', numeric_row_idx=1)
        d.save(p)
        before = p.read_bytes()
        with patch('guides.coursework_kfu_2025.table_continuation.repair_remaining_table_spills_inplace', return_value=0):
            assert refill_early_continuations_inplace(p) == 0
        assert p.read_bytes() == before
        print('PASS exact rollback on unresolved refill')
        d = fixture(12, 65)
        payload = _table_payload_rows(d)
        head = d.tables[0]
        tail = deepcopy(head._tbl)
        for row in list(tail.tr_lst):
            tail.remove(row)
        tail.append(deepcopy(head.rows[1]._tr))
        tail.append(head.rows[-1]._tr)
        head._tbl.addnext(tail)
        d.save(p)
        assert repair_remaining_table_spills_inplace(p) > 0
        lines = check_clean(p, payload)
        last_page = next(l.page_num for l in lines if 'Наблюдение номер 13' in l.text)
        assert any(l.page_num == last_page and 'Источник:' in l.text for l in lines)
        assert not any(l.text.strip() == '1 2' for l in lines if l.page_num > last_page)
        before = p.read_bytes()
        assert repair_remaining_table_spills_inplace(p) == 0 and p.read_bytes() == before
        print('PASS spilling head plus unlabelled final fragment measured as one table')
        # Authored text must never be consumed when looking for a continuation.
        d = fixture()
        tail = deepcopy(d.tables[0]._tbl)
        tail.remove(tail.tr_lst[0])
        d.tables[0]._tbl.addnext(tail)
        para = OxmlElement('w:p'); run = OxmlElement('w:r'); text = OxmlElement('w:t')
        text.text = 'Новый самостоятельный раздел'; run.append(text); para.append(run)
        d.tables[0]._tbl.addnext(para)
        assert _merge_adjacent_numeric_tail(d, 0) == 0
        print('PASS no merge across authored text')
        text.text = ''
        para.get_or_add_pPr().get_or_add_pageBreakBefore().val = True
        assert _merge_adjacent_numeric_tail(d, 0) == 0
        para.get_or_add_pPr().get_or_add_pageBreakBefore().val = False
        assert _merge_adjacent_numeric_tail(d, 0) == 1
        print('PASS authored page breaks protected; inactive breaks permit repair')
        d = fixture(20, 100, appendix=True)
        payload = _table_payload_rows(d)
        d.save(p)
        assert repair_remaining_table_spills_inplace(p) >= 2
        check_clean(p, payload)
        print('PASS captioned multi-page appendix is visible to spill detection')


if __name__ == '__main__':
    main()
