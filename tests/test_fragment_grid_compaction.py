"""Preserve physical cells while reconciling redundant continuation grids."""
import sys
from copy import deepcopy
from io import BytesIO
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from guides.coursework_kfu_2025.table_continuation import (
    apply_table_merging, _merge_continuation_chain_for_num,
)
from guides.coursework_kfu_2025.content_preservation import evaluate_content_preservation, table_cell_multiset
from guides.coursework_kfu_2025.rendered_table_validation import evaluate_table_layout_acceptance, build_rendered_table_identities


def fixture():
    d=Document();d.add_paragraph('Таблица 1.1.1')
    a=d.add_table(rows=2,cols=2)
    for r,values in zip(a.rows,[['Показатель','Описание'],['Начало','Первое значение']]):
        for c,v in zip(r.cells,values):c.text=v
    d.add_paragraph('Продолжение таблицы 1.1.1')
    b=d.add_table(rows=2,cols=2)
    for r,values in zip(b.rows,[['Показатель','Описание'],['Окончание','Последнее значение']]):
        for c,v in zip(r.cells,values):c.text=v
    col=OxmlElement('w:gridCol');col.set(qn('w:w'),'45');b._tbl.tblGrid.insert(1,col)
    b._tbl.tr_lst[0].tc_lst[0].grid_span=2
    b._tbl.tr_lst[1].tc_lst[1].grid_span=2
    return d


def clone(d):
    f=BytesIO();d.save(f);f.seek(0);return Document(f)


def main():
    for merge in [apply_table_merging,lambda d:_merge_continuation_chain_for_num(d,'1.1.1')]:
        d=fixture();source=clone(d)
        assert merge(d)>0
        assert len(d.tables)==1
        assert all(sum(c.grid_span for c in r.tc_lst)==2 for r in d.tables[0]._tbl.tr_lst)
        report,_=evaluate_content_preservation(source,d)
        assert not report['content_fail'],report
        before=d.element.xml;assert merge(d)==0 and d.element.xml==before
    print('PASS early/late staggered-grid merge, physical cells preserved, retry')
    d=fixture();d.tables[1]._tbl.tblGrid[1].set(qn('w:w'),'1500')
    before=d.element.xml
    assert _merge_continuation_chain_for_num(d,'1.1.1')==0 and before==d.element.xml
    assert apply_table_merging(d)==0 and before==d.element.xml
    print('PASS genuine incompatible grid retained')
    d=Document();t=d.add_table(rows=2,cols=2)
    for c in t.rows[0].cells:c.text='Заголовок'
    for c in t.rows[1].cells:c.text='Одинаковые данные'
    assert table_cell_multiset(d)['одинаковые данные']==2
    source=clone(d);t._tbl.remove(t.rows[1]._tr)
    report,_=evaluate_content_preservation(source,d)
    assert 'lost_table_cell_content' in report['content_fail']
    print('PASS equal text cells counted separately and real deletion caught')
    d=fixture();a,b=d.tables
    for row in list(b._tbl.tr_lst)[1:]:a._tbl.append(row)
    b._tbl.getparent().remove(b._tbl)
    blockers=evaluate_table_layout_acceptance([],build_rendered_table_identities(d),doc=d)
    assert any(b.blocker_type=='row_grid_extent_mismatch' for b in blockers)
    print('PASS over-wide physical row blocked')


if __name__=='__main__':main()
