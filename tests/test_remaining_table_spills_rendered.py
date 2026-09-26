"""Real LibreOffice regression: multi-page rows, repeat text and exact heights.
Run directly (requires the same soffice/pdfplumber dependencies as the worker).
"""
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from guides.coursework_kfu_2025.table_continuation import (
    repair_remaining_table_spills_inplace, _table_payload_rows,
    _cross_page_without_marker_probe,
)
from guides.coursework_kfu_2025.rendered_table_validation import (
    _cross_page_without_marker_blockers, build_rendered_table_identities,
)


def main():
    with tempfile.TemporaryDirectory() as temp:
        doc = Document()
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Cm(2)
        doc.styles['Normal'].font.size = Pt(12)
        doc.add_paragraph('Таблица 1.1.1')
        table = doc.add_table(rows=21, cols=2)
        table.style = 'Table Grid'
        table.cell(0, 0).text = 'Номер наблюдения'
        table.cell(0, 1).text = 'Описание наблюдения'
        for i, row in enumerate(table.rows[1:], 1):
            row.cells[0].text = f'Наблюдение номер {i}'
            row.cells[1].text = 'Повторяющееся описание\nВторая строка данных'
            row.height = Pt(100)
            row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        # A source row equal to the semantic header still has to be preserved
        # and must not become an unmarked header-only overflow page.
        table.cell(7, 0).text = table.cell(0, 0).text
        table.cell(7, 1).text = table.cell(0, 1).text
        doc.add_paragraph('Источник: составлено автором')
        source = Path(temp) / 'source.docx'
        target = Path(temp) / 'formatted.docx'
        doc.save(source)
        target.write_bytes(source.read_bytes())
        payload = _table_payload_rows(doc)
        repaired = repair_remaining_table_spills_inplace(target, source_docx_path=source)
        assert repaired >= 2, f'Expected several continuation pages, got {repaired}'
        result = Document(target)
        assert _table_payload_rows(result) == payload, 'Data changed or reordered'
        fails, _, lines = _cross_page_without_marker_probe(target, source)
        assert not fails, fails
        assert not _cross_page_without_marker_blockers(lines, build_rendered_table_identities(result))
        for part in result.tables[1:]:
            assert [c.text for c in part.rows[0].cells] == ['1', '2']
        before_retry = target.read_bytes()
        assert repair_remaining_table_spills_inplace(target, source_docx_path=source) == 0
        assert target.read_bytes() == before_retry, 'Retry mutated a clean document'
        print(f'PASS: {repaired} page splits, ordered content preserved, clean render, retry unchanged')
        doc.paragraphs[0].text = 'ПРИЛОЖЕНИЕ А'
        doc.save(source)
        target.write_bytes(source.read_bytes())
        repaired = repair_remaining_table_spills_inplace(target, source_docx_path=source)
        result = Document(target)
        assert repaired >= 2
        assert _table_payload_rows(result) == payload
        assert all(p.text == 'ПРОДОЛЖЕНИЕ ПРИЛОЖЕНИЯ А' for p in result.paragraphs
                   if p.text.startswith('ПРОДОЛЖЕНИЕ'))
        fails, _, lines = _cross_page_without_marker_probe(target, source)
        assert not fails, fails
        assert not _cross_page_without_marker_blockers(lines, build_rendered_table_identities(result))
        print(f'PASS: appendix, {repaired} page splits, correct labels and preserved rows')



if __name__ == '__main__':
    main()
