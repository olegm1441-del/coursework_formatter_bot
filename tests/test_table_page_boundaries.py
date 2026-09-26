"""Regression cases for page-local continuation evidence (no LO required)."""
import sys
import unittest
import tempfile
from unittest.mock import patch
from types import SimpleNamespace
from dataclasses import replace
from docx import Document
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from guides.coursework_kfu_2025.pdf_layout_analyzer import PdfLine
from guides.coursework_kfu_2025.rendered_table_validation import (
    RenderedTableIdentity, _cross_page_without_marker_blockers,
)

HEADER = 'Наименование показателя значение'
ROWS = ('первый показатель значение альфа', 'второй показатель значение бета',
        'третий показатель значение гамма')


def line(text, page, top):
    return PdfLine(text, page, top, top + 12)


def identity(caption='1.1.1', marker=None):
    return RenderedTableIdentity(0, 0, caption, marker, None,
                                 (HEADER,), None, (HEADER, *ROWS))


class PageBoundaryTests(unittest.TestCase):
    def test_next_caption_does_not_hide_tail_on_same_page(self):
        lines = [line('Таблица 1.1.1', 1, 100), line(ROWS[0], 1, 700),
                 line(ROWS[1], 2, 60), line('Таблица 1.1.2', 2, 300)]
        self.assertEqual(_cross_page_without_marker_blockers(lines, [identity()])[0]
                         .evidence['missing_marker_pages'], [2])

    def test_one_marker_does_not_cover_three_pages(self):
        lines = [line('Таблица 1.1.1', 1, 100), line(ROWS[0], 1, 200),
                 line('Продолжение таблицы 1.1.1', 2, 60), line(ROWS[1], 2, 100),
                 line(ROWS[2], 3, 60)]
        self.assertEqual(_cross_page_without_marker_blockers(lines, [identity()])[0]
                         .evidence['missing_marker_pages'], [3])

    def test_captionless_fragment_is_checked(self):
        lines = [line('Продолжение таблицы 1.1.1', 2, 60), line(ROWS[0], 2, 100),
                 line(ROWS[1], 3, 60)]
        self.assertTrue(_cross_page_without_marker_blockers(
            lines, [identity(None, 'Продолжение таблицы 1.1.1')]))

    def test_each_page_marked_is_valid(self):
        lines = [line('Таблица 1.1.1', 1, 100), line(ROWS[0], 1, 200),
                 line('Продолжение таблицы 1.1.1', 2, 60), line(ROWS[1], 2, 100),
                 line('Продолжение таблицы 1.1.1', 3, 60), line(ROWS[2], 3, 100)]
        self.assertFalse(_cross_page_without_marker_blockers(lines, [identity()]))

    def test_neighbour_or_prose_reusing_rows_is_not_our_table(self):
        for end in ('Таблица 1.1.2', 'Источник: составлено автором'):
            with self.subTest(end=end):
                lines = [line('Таблица 1.1.1', 1, 100), line(ROWS[0], 1, 200),
                         line(end, 1, 400), line(ROWS[1], 2, 60)]
                self.assertFalse(_cross_page_without_marker_blockers(lines, [identity()]))

    def test_similar_rows_in_next_fragment_do_not_extend_first(self):
        first = replace(identity(), following_marker='Продолжение таблицы 1.1.1')
        lines = [line('Таблица 1.1.1', 1, 100), line(ROWS[0], 1, 200),
                 line('Продолжение таблицы 1.1.1', 2, 60),
                 line(ROWS[0], 2, 100), line(ROWS[0], 3, 60)]
        self.assertFalse(_cross_page_without_marker_blockers(lines, [first]))

    def test_split_numeric_first_fragment_preserves_duplicate_rows(self):
        from guides.coursework_kfu_2025.table_continuation import (
            _split_cross_page_table_with_marker, _table_payload_rows)
        doc = Document()
        table = doc.add_table(rows=5, cols=2)
        for row, values in zip(table.rows, [('1', '2'), ('alpha', 'same'),
                                             ('alpha', 'same'), ('beta', 'value'),
                                             ('gamma', 'value')]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        before = _table_payload_rows(doc)
        self.assertTrue(_split_cross_page_table_with_marker(
            doc, 0, 2, '1.1.1', numeric_row_idx=0))
        self.assertEqual(before, _table_payload_rows(doc))
        self.assertEqual([c.text for c in doc.tables[1].rows[0].cells], ['1', '2'])
        self.assertEqual(len(doc.tables[0].rows), 3)

    def test_instrumentation_prefix_keeps_source_unchanged(self):
        from guides.coursework_kfu_2025.table_markers import instrument_table_rows_copy
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / 'source.docx'
            doc = Document()
            doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Длинное слово'
            doc.save(source)
            before = source.read_bytes()
            result = instrument_table_rows_copy(source, 0, workdir=temp)
            text = Document(result.instrumented_docx_path).tables[0].cell(0, 0).text
            self.assertTrue(text.startswith(result.row_markers[0]))
            self.assertTrue(text.endswith('Длинное слово'))
            self.assertEqual(source.read_bytes(), before)

    def test_appendix_does_not_match_reused_rows_in_body(self):
        item = replace(identity(None), appendix_anchor='ПРИЛОЖЕНИЕ А')
        lines = [line(ROWS[0], 46, 100), line('ПРИЛОЖЕНИЕ А', 63, 60),
                 line(ROWS[0], 63, 200), line(ROWS[1], 64, 80)]
        blocker = _cross_page_without_marker_blockers(lines, [item])[0]
        self.assertEqual(blocker.table_num, 'appendix:А')
        self.assertEqual(blocker.evidence['data_pages'], [63, 64])

    def test_appendix_split_uses_appendix_marker(self):
        from guides.coursework_kfu_2025.table_continuation import _split_cross_page_table_with_marker
        doc = Document()
        table = doc.add_table(rows=4, cols=2)
        for i, row in enumerate(table.rows):
            for cell in row.cells:
                cell.text = f'row {i}'
        self.assertTrue(_split_cross_page_table_with_marker(
            doc, 0, 2, 'appendix:А', numeric_row_idx=None))
        self.assertEqual([p.text for p in doc.paragraphs], ['ПРОДОЛЖЕНИЕ ПРИЛОЖЕНИЯ А'])

    def test_remaining_spill_rolls_back_on_new_layout_failure(self):
        from guides.coursework_kfu_2025 import table_continuation as tc
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'input.docx'
            doc = Document()
            doc.add_paragraph('Таблица 1.1.1')
            table = doc.add_table(rows=4, cols=2)
            for i, row in enumerate(table.rows):
                for j, cell in enumerate(row.cells):
                    cell.text = f'row {i} column {j}'
            doc.save(path)
            before = path.read_bytes()
            spill = SimpleNamespace(table_num='1.1.1', evidence={'table_index': 0})
            baseline = {('single_table_crosses_pages_without_marker', '1.1.1')}
            broken = {('orphaned_header_row', '2.1.1')}
            with patch.object(tc, '_cross_page_without_marker_probe', side_effect=[
                    (baseline, {'1.1.1'}, []), (broken, set(), []),
                    (broken, set(), []), (baseline, {'1.1.1'}, [])]), \
                 patch('guides.coursework_kfu_2025.rendered_table_validation.'
                       '_cross_page_without_marker_blockers', side_effect=[
                           [spill], [spill], [], [spill]]), \
                 patch.object(tc, '_instrumented_data_row_pages',
                              side_effect=[{1: 1, 2: 1, 3: 2}, {2: 1, 3: 1}]):
                self.assertEqual(tc.repair_remaining_table_spills_inplace(path), 0)
            self.assertEqual(path.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
