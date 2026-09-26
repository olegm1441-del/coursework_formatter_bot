"""Render regressions for balanced placement after a whole-table move."""
import sys,tempfile
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from guides.coursework_kfu_2025.table_continuation import (
    refill_table_page_gaps_inplace, repair_table_adjacency_inplace, keep_short_tables_whole_inplace, _table_payload_rows,
    _cross_page_without_marker_probe,
)


def fixture(rows=5,height=80,space=300):
    d=Document();s=d.sections[0];s.page_height=Cm(29.7)
    s.top_margin=s.bottom_margin=Cm(2)
    d.styles['Normal'].font.size=Pt(12)
    d.add_paragraph('Предшествующий текст').paragraph_format.space_after=Pt(space)
    cap=d.add_paragraph('Таблица 1.1.1')
    cap.paragraph_format.page_break_before=True
    cap.paragraph_format.keep_with_next=True
    t=d.add_table(rows=rows+1,cols=2);t.style='Table Grid'
    for i,r in enumerate(t.rows):
        r.cells[0].text='Показатель' if i==0 else f'Показатель номер {i}'
        r.cells[1].text='Значение' if i==0 else f'Исходное значение {i}'
        r.height=Pt(25 if i==0 else height);r.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in r.cells:c.paragraphs[0].paragraph_format.keep_with_next=i<rows
    d.add_paragraph('Источник: исходные данные')
    d.add_paragraph('Примечание: исходное пояснение')
    return d


def main():
    with tempfile.TemporaryDirectory() as tmp:
        p=Path(tmp)/'trial.docx'
        d=fixture();payload=_table_payload_rows(d);d.save(p)
        assert refill_table_page_gaps_inplace(p)==1
        result=Document(p)
        assert len(result.tables)==2
        assert _table_payload_rows(result)==payload
        fails,_,lines=_cross_page_without_marker_probe(p,None)
        assert not fails,fails
        assert any(l.page_num==1 and 'Таблица 1.1.1'==l.text.strip() for l in lines)
        for text in ['Продолжение таблицы 1.1.1','Показатель номер 5','Источник: исходные данные','Примечание: исходное пояснение']:
            assert any(l.page_num==2 and text in l.text for l in lines),text
        before=p.read_bytes();assert refill_table_page_gaps_inplace(p)==0
        assert p.read_bytes()==before
        print('PASS fill preceding page, continuation and notes together, preserve rows, retry')
        for d in [fixture(12,45,430),fixture(5,80,600)]:
            d.save(p);before=p.read_bytes()
            assert refill_table_page_gaps_inplace(p)==0
            assert p.read_bytes()==before
        print('PASS retain whole table when little data fits or page space is insufficient')
        d=fixture();d.save(p);before=p.read_bytes()
        with patch('guides.coursework_kfu_2025.table_continuation.repair_remaining_table_spills_inplace',return_value=0):
            assert refill_table_page_gaps_inplace(p)==0
        assert p.read_bytes()==before
        print('PASS rollback when a continuation cannot be verified')
        d=fixture();d.paragraphs[1].text='ПРИЛОЖЕНИЕ А';d.save(p);before=p.read_bytes()
        assert refill_table_page_gaps_inplace(p)==0 and p.read_bytes()==before
        print('PASS appendix untouched')
        d=fixture(3,80,440)
        d.paragraphs[1].paragraph_format.page_break_before=False
        for row in d.tables[0].rows:
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.keep_with_next=False
        payload=_table_payload_rows(d);d.save(p)
        assert repair_table_adjacency_inplace(p)>0
        assert _table_payload_rows(Document(p))==payload
        fails,_,lines=_cross_page_without_marker_probe(p,None)
        assert not fails,fails
        last_page=next(l.page_num for l in lines if 'Показатель номер 3' in l.text)
        for text in ['Источник: исходные данные','Примечание: исходное пояснение']:
            assert any(l.page_num==last_page and text in l.text for l in lines),text
        before=p.read_bytes();assert repair_table_adjacency_inplace(p)==0
        assert p.read_bytes()==before
        for para in Document(p).paragraphs:
            if para.text.startswith(('Источник:', 'Примечание:')):
                assert para.paragraph_format.keep_together is True
        print('PASS detached notes move with final data row, whole paragraphs and retry')
        # A renderer can clip a tail without reporting a second table page.
        # Simulate that initial observation; the accepted trial still needs a
        # real render with all original data rows visible on the fresh page.
        from guides.coursework_kfu_2025.pdf_layout_analyzer import PdfLine
        d=fixture(3,40,430)
        d.paragraphs[1].paragraph_format.page_break_before=False
        d.save(p);payload=_table_payload_rows(d)
        real_probe=_cross_page_without_marker_probe
        calls=0
        def clipped_once(*args,**kwargs):
            nonlocal calls
            calls+=1
            if calls==1:
                return set(),set(),[
                    PdfLine('Таблица 1.1.1',1,510,525),
                    PdfLine('Показатель номер 1',1,540,555)]
            return real_probe(*args,**kwargs)
        with patch('guides.coursework_kfu_2025.table_continuation._cross_page_without_marker_probe',side_effect=clipped_once), patch('guides.coursework_kfu_2025.table_continuation._instrumented_data_row_pages',return_value=None):
            assert keep_short_tables_whole_inplace(p)==1
        assert _table_payload_rows(Document(p))==payload
        _,_,lines=real_probe(p,None)
        for number in (1,2,3):
            assert any(l.page_num==2 and f'Показатель номер {number}' in l.text for l in lines)
        print('PASS clipped tail detected without cross-page warning; all rows verified after move')
        from guides.coursework_kfu_2025.table_markers import map_table_rows_to_pages
        d=Document();t=d.add_table(rows=4,cols=2);t.autofit=False
        t.columns[0].width=Pt(12);t.columns[1].width=Pt(220)
        for i,row in enumerate(t.rows):
            row.cells[0].width=Pt(0);row.cells[1].width=Pt(0)
            row.cells[0].text='Год' if i==0 else str(i)
            row.cells[1].text='Комментарий' if i==0 else f'Исходное описание этапа {i}'
        d.save(p);before=p.read_bytes()
        mapped=map_table_rows_to_pages(p,0,allow_repeated_header=True)
        assert mapped.row_pages=={0:1,1:1,2:1,3:1},mapped
        assert not mapped.missing_rows and p.read_bytes()==before
        print('PASS narrow year cells do not conceal diagnostic row markers')
        d=fixture();d.paragraphs[1].insert_paragraph_before('')
        d.save(p);assert repair_table_adjacency_inplace(p)>0
        result=Document(p)
        assert not any(not para.text for para in result.paragraphs)
        print('PASS empty spacer removed before forced table caption')


if __name__=='__main__':main()
