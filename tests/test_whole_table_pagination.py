"""Rendered regressions: short tables stay whole, long tables remain splittable."""
import sys
import tempfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from guides.coursework_kfu_2025.table_continuation import (
    keep_short_tables_whole_inplace, _table_payload_rows,
    keep_short_uncaptioned_tables_whole_inplace,
    _cross_page_without_marker_probe, _split_cross_page_table_with_marker,
    repair_same_page_appendix_continuations_inplace,
    fit_oversized_table_grids,
)
from guides.coursework_kfu_2025.rendered_table_validation import (
    _appendix_label_blockers, build_rendered_table_identities,
)
from guides.coursework_kfu_2025.pdf_layout_analyzer import PdfLine
from guides.coursework_kfu_2025.page_breaks import apply_page_breaks
from guides.coursework_kfu_2025.contents_builder import _reapply_front_matter_layout


def fixture(rows=8):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    doc.styles['Normal'].font.size = Pt(12)
    p = doc.add_paragraph('Предшествующий текст')
    p.paragraph_format.space_after = Pt(430)
    doc.add_paragraph('Таблица 1.1.1').paragraph_format.keep_with_next = True
    table = doc.add_table(rows=rows + 1, cols=2)
    table.style = 'Table Grid'
    for i, row in enumerate(table.rows):
        row.cells[0].text = 'Название' if i == 0 else f'Показатель номер {i}'
        row.cells[1].text = 'Описание' if i == 0 else f'Значение показателя {i}'
        row.height = Pt(40)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    doc.add_paragraph('Источник: составлено автором')
    return doc


def test_appendices():
    for blank in ('', None):
        doc = Document()
        for text in ['Титульный лист', 'СОДЕРЖАНИЕ', 'ВВЕДЕНИЕ', 'Текст', 'ПРИЛОЖЕНИЯ']:
            doc.add_paragraph(text)
        if blank is not None:
            doc.add_paragraph(blank)
        first = doc.add_paragraph('ПРИЛОЖЕНИЕ А')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Данные'
        second = doc.add_paragraph('ПРИЛОЖЕНИЕ Б')
        apply_page_breaks(doc, 2)
        assert first.paragraph_format.page_break_before is False
        assert second.paragraph_format.page_break_before is True
        _reapply_front_matter_layout(doc)
        assert first.paragraph_format.page_break_before is False, 'TOC rebuilt orphan appendix page'
        assert second.paragraph_format.page_break_before is True
    lines = [PdfLine('ПРИЛОЖЕНИЯ', 4, 60, 75), PdfLine('4', 4, 800, 814),
             PdfLine('ПРИЛОЖЕНИЕ А', 5, 60, 75)]
    assert _appendix_label_blockers(lines)[0].blocker_type == 'appendix_heading_isolated'
    lines[-1] = PdfLine('ПРИЛОЖЕНИЕ А', 4, 90, 105)
    assert not _appendix_label_blockers(lines)
    doc = Document()
    doc.add_paragraph('ПРИЛОЖЕНИЯ\t54')
    doc.add_paragraph('ВВЕДЕНИЕ')
    doc.add_paragraph('Таблица 1.1.1')
    doc.add_table(rows=2, cols=2)
    assert build_rendered_table_identities(doc)[0].appendix_anchor is None
    print('PASS appendix heading and first label share page, including TOC rebuild')


def main():
    test_appendices()
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / 'uncaptioned.docx'
        for count, expected in ((8, 1), (30, 0)):
            doc = fixture(count)
            caption = doc.paragraphs[1]._p
            caption.getparent().remove(caption)
            doc.save(path)
            payload = _table_payload_rows(doc)
            assert keep_short_uncaptioned_tables_whole_inplace(path) == expected
            assert _table_payload_rows(Document(path)) == payload
            before = path.read_bytes()
            assert keep_short_uncaptioned_tables_whole_inplace(path) == 0
            assert path.read_bytes() == before
        print('PASS uncaptioned short/long table, content preservation and retry')
    with tempfile.TemporaryDirectory() as temp:
        p = Path(temp) / 'trial.docx'
        doc = Document()
        table = doc.add_table(rows=2, cols=3)
        table.style = 'Table Grid'
        for col in table.columns:
            col.width = Cm(9)
        table.cell(0, 0).merge(table.cell(0, 1)).text = 'Объединённая ячейка'
        table.cell(1, 2).text = 'Крайняя правая ячейка'
        landscape = doc.add_section(WD_SECTION.NEW_PAGE)
        landscape.page_width = Cm(29.7); landscape.page_height = Cm(21)
        landscape.left_margin = landscape.right_margin = Cm(2)
        wide = doc.add_table(rows=2, cols=3)
        for col in wide.columns:
            col.width = Cm(8)
        wide_xml = wide._tbl.xml
        payload = _table_payload_rows(doc)
        assert fit_oversized_table_grids(doc) == 1
        assert _table_payload_rows(doc) == payload
        assert wide._tbl.xml == wide_xml, 'valid landscape table changed'
        widths = [int(c.get(qn('w:w'))) for c in table._tbl.tblGrid]
        merged_width = int(table.cell(0, 0)._tc.tcPr.find(qn('w:tcW')).get(qn('w:w')))
        assert merged_width == sum(widths[:2]), 'merged cell width corrupted'
        before_xml = doc.element.xml
        assert fit_oversized_table_grids(doc) == 0
        assert doc.element.xml == before_xml
        print('PASS oversized grid fits section, merged cells and landscape preserved')
        for split in (False, True):
            doc = fixture()
            if split:
                # A numeric-looking source data row must survive a merge.
                doc.tables[0].cell(6, 0).text = '1'
                doc.tables[0].cell(6, 1).text = '2'
                assert _split_cross_page_table_with_marker(doc, 0, 3, '1.1.1', numeric_row_idx=None)
            payload = _table_payload_rows(doc)
            doc.save(p)
            assert keep_short_tables_whole_inplace(p) == 1, f'whole table, split={split}'
            result = Document(p)
            assert len(result.tables) == 1
            assert _table_payload_rows(result) == payload
            if split:
                assert sum([c.text for c in row.cells] == ['1', '2']
                           for row in result.tables[0].rows) == 2
            fails, _, lines = _cross_page_without_marker_probe(p, None)
            # The legacy validator flags the source's numeric data row too;
            # keeping that source warning is preferable to deleting the data.
            expected = {('same_page_numeric_continuation', '1.1.1')} if split else set()
            assert fails == expected, fails
            assert not any('Продолжение таблицы' in line.text for line in lines)
            before = p.read_bytes()
            assert keep_short_tables_whole_inplace(p) == 0
            assert p.read_bytes() == before
            print(f'PASS short table on one page, split={split}, content and retry preserved')
        doc = fixture(22)
        doc.save(p)
        before = p.read_bytes()
        assert keep_short_tables_whole_inplace(p) == 0
        assert p.read_bytes() == before, 'long table was mutated'
        print('PASS long table not forced onto one page')
        # Underestimate a long table deliberately: only the real render may accept.
        with patch('guides.coursework_kfu_2025.table_continuation._estimate_row_height', return_value=1):
            assert keep_short_tables_whole_inplace(p) == 0
        assert p.read_bytes() == before, 'failed render trial did not restore exact bytes'
        print('PASS misleading height estimate cannot authorize merge')
        for height in (80, 430):
            doc = Document()
            doc.add_paragraph('ПРИЛОЖЕНИЕ 1')
            first = doc.add_table(rows=2, cols=2)
            first.style = 'Table Grid'
            for row in first.rows:
                row.cells[0].text = 'Исходное наблюдение'
                row.cells[1].text = 'Исходное значение'
            first.rows[1].height = Pt(240)
            first.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            doc.add_paragraph('ПРОДОЛЖЕНИЕ ПРИЛОЖЕНИЯ 1')
            second = doc.add_table(rows=2, cols=2)
            second.style = 'Table Grid'
            second.cell(0, 0).text = '1'; second.cell(0, 1).text = '2'
            second.cell(1, 0).text = 'Последующее наблюдение'
            second.cell(1, 1).text = 'Последующее значение'
            second.rows[1].height = Pt(height)
            second.rows[1].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            for row in second.rows:
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            doc.add_paragraph('Источник: составлено автором')
            payload = _table_payload_rows(doc)
            doc.save(p)
            assert repair_same_page_appendix_continuations_inplace(p) == 1
            result = Document(p)
            assert _table_payload_rows(result) == payload
            assert len(result.tables) == (1 if height == 80 else 2)
            fails, _, lines = _cross_page_without_marker_probe(p, None)
            assert not fails, fails
            before = p.read_bytes()
            assert repair_same_page_appendix_continuations_inplace(p) == 0
            assert before == p.read_bytes()
            print(f'PASS appendix continuation cleanup, height={height}, rows and retry preserved')


if __name__ == '__main__':
    main()
